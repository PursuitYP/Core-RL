#!/usr/bin/env python3
"""Build the paper-structured two-column poster from the supplied template."""

from __future__ import annotations

import hashlib
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "poster_template.pptx"
OUTPUT = ROOT / "poster_final.pptx"
ADAPTER_ARCHIVE = (
    ROOT
    / "所有的运行结果"
    / "02_From_Predictive_Feature_Geometry_to_Control_Utilization__integrated_findings.zip"
)
ADAPTER_FIGURE = "02_integrated_paper/figures/fig02_stage_b1_trace_only_rescue.png"


def color(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


NAVY = color("17324D")
BLUE = color("3E718F")
GREEN = color("477F6B")
OCHRE = color("94652F")
INK = color("27343C")
MID = color("5C6A72")
GRAY = color("6F7D84")
LIGHT_LINE = color("D8E0E4")
PALE_BLUE = color("ECF2F5")
PALE_GREEN = color("EDF4F1")
SOFT = color("F7F8F8")
WHITE = color("FFFFFF")
FONT = "Arial"


def inches(value: float):
    return Inches(value)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clear_slide(slide) -> None:
    tree = slide.shapes._spTree
    for shape in list(slide.shapes):
        tree.remove(shape._element)


def name_shape(shape, name: str | None):
    if name:
        shape.name = name
    return shape


def set_shape_text(
    shape,
    text: str,
    size: float,
    foreground: RGBColor = INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.02,
    line_spacing: float = 1.0,
) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = valign
    frame.margin_left = inches(margin)
    frame.margin_right = inches(margin)
    frame.margin_top = inches(margin)
    frame.margin_bottom = inches(margin)
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = line_spacing
        paragraph.space_after = Pt(0)
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = foreground


def add_text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    size: float,
    foreground: RGBColor = INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.02,
    line_spacing: float = 1.0,
    name: str | None = None,
):
    shape = slide.shapes.add_textbox(inches(x), inches(y), inches(w), inches(h))
    set_shape_text(
        shape,
        text,
        size,
        foreground,
        bold,
        align,
        valign,
        margin,
        line_spacing,
    )
    return name_shape(shape, name)


def add_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: RGBColor = WHITE,
    line: RGBColor = LIGHT_LINE,
    line_width: float = 0.8,
    name: str | None = None,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        inches(x),
        inches(y),
        inches(w),
        inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(line_width)
    return name_shape(shape, name)


def add_rule(
    slide,
    x: float,
    y: float,
    w: float,
    foreground: RGBColor = LIGHT_LINE,
    height: float = 0.025,
    name: str | None = None,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        inches(x),
        inches(y),
        inches(w),
        inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = foreground
    shape.line.fill.background()
    return name_shape(shape, name)


def add_vertical_rule(
    slide,
    x: float,
    y: float,
    h: float,
    foreground: RGBColor = LIGHT_LINE,
    width: float = 0.025,
    name: str | None = None,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        inches(x),
        inches(y),
        inches(width),
        inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = foreground
    shape.line.fill.background()
    return name_shape(shape, name)


def add_label(slide, x: float, y: float, w: float, text: str, accent: RGBColor):
    return add_text(
        slide,
        x,
        y,
        w,
        0.34,
        text.upper(),
        14.4,
        accent,
        True,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_tmaze(slide, x: float, y: float) -> None:
    """Compact running example; all branches stay outside text boxes."""
    cue = add_rect(
        slide,
        x,
        y + 0.34,
        1.00,
        0.60,
        fill=PALE_BLUE,
        line=BLUE,
        line_width=0.9,
        name="T-maze cue",
    )
    set_shape_text(
        cue,
        "cue L/R",
        14.0,
        BLUE,
        True,
        PP_ALIGN.CENTER,
        MSO_ANCHOR.MIDDLE,
        0,
    )
    cell_x = x + 1.22
    for index in range(5):
        cell = add_rect(
            slide,
            cell_x + 0.61 * index,
            y + 0.34,
            0.50,
            0.60,
            fill=SOFT,
            line=color("BAC6CC"),
            line_width=0.65,
            name=f"Aliased cell {index + 1}",
        )
        set_shape_text(
            cell,
            "?",
            15.0,
            GRAY,
            True,
            PP_ALIGN.CENTER,
            MSO_ANCHOR.MIDDLE,
            0,
        )
    junction_x = cell_x + 3.08
    junction = add_rect(
        slide,
        junction_x,
        y + 0.34,
        0.56,
        0.60,
        fill=PALE_GREEN,
        line=GREEN,
        line_width=0.9,
        name="T-maze junction",
    )
    set_shape_text(
        junction,
        "J",
        15.0,
        GREEN,
        True,
        PP_ALIGN.CENTER,
        MSO_ANCHOR.MIDDLE,
        0,
    )
    branch_x = junction_x + 0.27
    add_vertical_rule(slide, branch_x, y, 0.34, GREEN, 0.025)
    add_vertical_rule(slide, branch_x, y + 0.96, 0.25, GREEN, 0.025)
    add_text(
        slide,
        junction_x + 0.66,
        y - 0.03,
        0.72,
        0.28,
        "turn L",
        13.2,
        GREEN,
        True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        junction_x + 0.66,
        y + 0.98,
        0.72,
        0.25,
        "turn R",
        13.2,
        GREEN,
        True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        cell_x,
        y + 0.96,
        2.95,
        0.25,
        "aliased corridor",
        13.2,
        MID,
        False,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_gate_chain(slide, x: float, y: float) -> None:
    gates = [
        ("PRESERVE", "retain the hidden\naction distinction", BLUE),
        ("EXPOSE", "make it learnable by\nthe online controller", OCHRE),
        ("ADAPT", "keep a compact basis\nas the world drifts", GREEN),
    ]
    starts = [x, x + 3.02, x + 6.04]
    for start, (label, detail, accent) in zip(starts, gates):
        add_text(
            slide,
            start,
            y,
            2.35,
            0.32,
            label,
            14.2,
            accent,
            True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
        add_text(
            slide,
            start,
            y + 0.36,
            2.35,
            0.62,
            detail,
            14.2,
            INK,
            False,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.TOP,
            line_spacing=0.94,
        )
    for arrow_x in (x + 2.49, x + 5.51):
        add_text(
            slide,
            arrow_x,
            y + 0.24,
            0.36,
            0.44,
            "→",
            20.0,
            GRAY,
            False,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )


def add_study_row(
    slide,
    x: float,
    y: float,
    w: float,
    rq: str,
    title: str,
    design: str,
    accent: RGBColor,
) -> None:
    add_text(
        slide,
        x,
        y,
        0.62,
        0.40,
        rq,
        14.2,
        accent,
        True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        x + 0.70,
        y - 0.02,
        w - 0.70,
        0.42,
        title,
        16.5,
        NAVY,
        True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        x + 0.70,
        y + 0.43,
        w - 0.74,
        0.76,
        design,
        15.2,
        INK,
        False,
        valign=MSO_ANCHOR.TOP,
        line_spacing=0.96,
    )
    add_rule(slide, x, y + 1.22, w, LIGHT_LINE)


def add_bank_mechanism(slide, x: float, y: float, w: float) -> None:
    """Explain the paper's online GVF-bank loop without adding card clutter."""
    stages = [
        (
            "1  GENERATE",
            "memory · timing\ncontrol · regime\n+ sparse mixtures",
            BLUE,
        ),
        (
            "2  SCORE",
            "decision / decode gains\n+ uncertainty reduction\n− TD error / redundancy",
            OCHRE,
        ),
        (
            "3  RETAIN + EXPOSE",
            "keep k complementary GVFs\n[oₜ, bank, 1]\n→ streaming linear policy",
            GREEN,
        ),
    ]
    gap = 0.48
    stage_w = (w - 2 * gap) / 3
    for index, (heading, detail, accent) in enumerate(stages):
        stage_x = x + index * (stage_w + gap)
        add_rule(slide, stage_x, y, stage_w, accent, 0.035)
        add_text(
            slide,
            stage_x,
            y + 0.13,
            stage_w,
            0.32,
            heading,
            13.8,
            accent,
            True,
            valign=MSO_ANCHOR.MIDDLE,
        )
        add_text(
            slide,
            stage_x,
            y + 0.49,
            stage_w,
            0.88,
            detail,
            13.8,
            INK,
            False,
            line_spacing=0.93,
        )
        if index < 2:
            add_text(
                slide,
                stage_x + stage_w + 0.05,
                y + 0.55,
                gap - 0.10,
                0.38,
                "→",
                18.0,
                GRAY,
                False,
                align=PP_ALIGN.CENTER,
                valign=MSO_ANCHOR.MIDDLE,
            )


def add_geometry_result_table(slide, x: float, y: float, w: float) -> None:
    columns = [2.05, 3.35, w - 5.40]
    headers = ["Environment", "Selected result", "What it rules out"]
    starts = [x, x + columns[0], x + columns[0] + columns[1]]
    add_rect(slide, x, y, w, 0.40, fill=SOFT, line=SOFT, line_width=0)
    for start, width, header in zip(starts, columns, headers):
        add_text(
            slide,
            start + 0.10,
            y + 0.04,
            width - 0.18,
            0.31,
            header,
            13.8,
            MID,
            True,
            valign=MSO_ANCHOR.MIDDLE,
        )
    rows = [
        (
            "T-maze",
            "non-oracle ≈ 0.50 (chance)",
            "feature shape alone does not recover memory control",
        ),
        (
            "Two-loop",
            "moment 0.797±0.068\nraw 0.495±0.004",
            "a positive effect can be environment-specific",
        ),
        (
            "Hidden velocity",
            "raw −0.063±0.018\nhighest mean",
            "regularization and priors can reduce control utility",
        ),
    ]
    row_h = 0.57
    for index, row in enumerate(rows):
        row_y = y + 0.42 + index * row_h
        if index:
            add_rule(slide, x, row_y, w, color("E8ECEE"), 0.018)
        for start, width, value in zip(starts, columns, row):
            add_text(
                slide,
                start + 0.10,
                row_y + 0.05,
                width - 0.18,
                row_h - 0.08,
                value,
                14.0,
                INK,
                index == 1 and start == starts[1],
                valign=MSO_ANCHOR.MIDDLE,
                line_spacing=0.92,
            )


def add_bank_table(slide, x: float, y: float, w: float) -> None:
    columns = [5.85, 2.75, w - 8.60]
    headers = ["Bank variant", "Cue decodability", "Post-switch control"]
    starts = [x, x + columns[0], x + columns[0] + columns[1]]
    add_rect(slide, x, y, w, 0.40, fill=SOFT, line=SOFT, line_width=0)
    for start, width, header in zip(starts, columns, headers):
        add_text(
            slide,
            start + 0.10,
            y + 0.04,
            width - 0.18,
            0.31,
            header,
            13.8,
            MID,
            True,
            align=PP_ALIGN.LEFT if start == starts[0] else PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
    rows = [
        ("k=2 full", "0.6483", "0.4950"),
        ("k=3 full", "0.7294", "0.5565"),
        ("k=4 full", "0.7386", "0.5408"),
        ("k=3, no diversity", "0.6951", "0.5196"),
        ("k=3, no novelty", "0.7294", "0.5565"),
    ]
    row_h = 0.39
    for index, row in enumerate(rows):
        row_y = y + 0.41 + index * row_h
        if index in (1, 4):
            add_rect(
                slide,
                x,
                row_y,
                w,
                row_h,
                fill=PALE_GREEN,
                line=PALE_GREEN,
                line_width=0,
                name=f"Highlighted k=3 row {index}",
            )
        elif index:
            add_rule(slide, x, row_y, w, color("E8ECEE"), 0.016)
        for start, width, value in zip(starts, columns, row):
            add_text(
                slide,
                start + 0.10,
                row_y + 0.03,
                width - 0.18,
                row_h - 0.05,
                value,
                14.0,
                GREEN if index in (1, 4) else INK,
                index in (1, 4),
                align=PP_ALIGN.LEFT if start == starts[0] else PP_ALIGN.CENTER,
                valign=MSO_ANCHOR.MIDDLE,
            )


def build_poster() -> None:
    template_hash = sha256(TEMPLATE)
    presentation = Presentation(str(TEMPLATE))
    presentation.core_properties.title = (
        "Exploring Useful Predictive Representations for Streaming RL under Partial "
        "Observability and Non-Stationarity"
    )
    presentation.core_properties.author = "Mingzhu Li; Xinwei Song; Yuqi Wei; Peng Yu"
    presentation.core_properties.subject = "2026 RL course project poster"
    presentation.core_properties.keywords = (
        "streaming reinforcement learning, partial observability, GVF, predictive state, "
        "control accessibility, non-stationarity"
    )
    if len(presentation.slides) != 1:
        raise RuntimeError("Expected a one-slide poster template")

    slide = presentation.slides[0]
    clear_slide(slide)
    slide_w = presentation.slide_width / 914400
    slide_h = presentation.slide_height / 914400

    # Keep both institutional color bands. Only neutralize the central reading area.
    add_rect(
        slide,
        0,
        6.46,
        slide_w,
        32.58 - 6.46,
        fill=WHITE,
        line=WHITE,
        line_width=0,
        name="Neutral body",
    )

    add_text(
        slide,
        0.92,
        1.76,
        21.78,
        2.82,
        "Exploring Useful Predictive Representations for Streaming RL\nunder Partial Observability and Non-Stationarity",
        51.0,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        line_spacing=0.94,
        name="Poster title",
    )
    add_text(
        slide,
        1.20,
        4.82,
        21.22,
        0.66,
        "Mingzhu Li   ·   Xinwei Song   ·   Yuqi Wei   ·   Peng Yu",
        31.0,
        WHITE,
        False,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        name="Authors",
    )

    # Full-width opening.
    add_label(slide, 0.92, 6.76, 3.20, "Problem statement", BLUE)
    add_text(
        slide,
        0.92,
        7.16,
        21.78,
        0.72,
        "When does predictive knowledge become a useful state representation for replay-free streaming control?",
        27.0,
        NAVY,
        True,
        valign=MSO_ANCHOR.MIDDLE,
        name="Core question",
    )
    add_text(
        slide,
        0.92,
        7.97,
        21.78,
        0.78,
        "We test a three-stage bottleneck: useful predictive state must preserve action-relevant distinctions, expose them to the learner, and remain compact and complementary under drift.",
        17.0,
        MID,
        False,
        valign=MSO_ANCHOR.MIDDLE,
        line_spacing=0.98,
        name="Core hypothesis",
    )
    add_rule(slide, 0.92, 9.01, 21.78, LIGHT_LINE, 0.035)

    left_x, left_w = 0.92, 8.92
    right_x, right_w = 10.64, 12.06
    add_vertical_rule(
        slide,
        10.22,
        9.34,
        20.62,
        color("E4E9EC"),
        0.024,
        name="Column divider",
    )

    # ------------------------------------------------------------------
    # LEFT COLUMN — why the problem matters and how the experiments isolate it.
    # ------------------------------------------------------------------
    add_label(slide, left_x, 9.34, left_w, "Why this problem?", BLUE)
    add_text(
        slide,
        left_x,
        9.72,
        left_w,
        0.58,
        "Accurate prediction ≠ useful control state",
        22.5,
        NAVY,
        True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        left_x,
        10.34,
        left_w,
        1.20,
        "The Alberta Plan proposes predictions as agent state. In a partially observable Markov\n"
        "decision process (POMDP), retained hidden information may still be unusable by the\n"
        "online controller. We therefore separate three failure modes—preservation,\n"
        "accessibility, and adaptation under drift—rather than rank predictions alone.",
        16.6,
        INK,
        False,
        line_spacing=1.00,
    )
    add_gate_chain(slide, left_x + 0.08, 11.70)
    add_rule(slide, left_x, 12.82, left_w, LIGHT_LINE)

    add_label(slide, left_x, 13.08, left_w, "Running example", GREEN)
    add_text(
        slide,
        left_x,
        13.45,
        left_w,
        0.48,
        "Continuing delayed-cue T-maze",
        19.5,
        NAVY,
        True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_tmaze(slide, left_x + 0.08, 13.97)
    add_text(
        slide,
        left_x + 6.10,
        13.99,
        2.72,
        1.18,
        "Cue appears once.\nCorridor observations repeat.\nThe junction requires the past cue.",
        14.5,
        INK,
        False,
        line_spacing=0.94,
    )
    add_text(
        slide,
        left_x + 0.08,
        15.25,
        left_w - 0.16,
        0.53,
        "RL task: cue → aliased corridor → junction turn; correct / incorrect feedback\n"
        "Goal: continuing accuracy · Drift: cue corruption · distractors · timing · remapping",
        13.3,
        MID,
        False,
        valign=MSO_ANCHOR.TOP,
        line_spacing=0.92,
    )
    add_rule(slide, left_x, 15.81, left_w, LIGHT_LINE)

    add_label(slide, left_x, 16.07, left_w, "How we test the claim", BLUE)
    add_text(
        slide,
        left_x,
        16.43,
        left_w,
        0.48,
        "xₜ = [ oₜ , f(gₜ) , 1 ]",
        20.0,
        NAVY,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        left_x,
        16.92,
        left_w,
        0.57,
        "gₜ: learned forecasts from general value functions (GVFs)—cue, timing, regime\n"
        "f: identity, residual random Fourier features (RFF), tile coding, or bank interface",
        13.4,
        MID,
        False,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.TOP,
        line_spacing=0.90,
    )
    add_study_row(
        slide,
        left_x,
        17.57,
        left_w,
        "RQ1",
        "Representation geometry",
        "4 POMDPs · 10 conditions · 3 seeds; compare transforms within environment.\n"
        "Accuracy and reward differ across tasks, so results are never pooled.",
        BLUE,
    )
    add_study_row(
        slide,
        left_x,
        19.06,
        left_w,
        "RQ2",
        "Control accessibility",
        "Trace-only T-maze · 10 paired seeds; hold gₜ and the learner fixed,\n"
        "then vary only the controller adapter: identity, RFF, or tile coding.",
        OCHRE,
    )
    add_study_row(
        slide,
        left_x,
        20.55,
        left_w,
        "RQ3",
        "Bank structure under drift",
        "Big-world T-maze · 4 seeds; vary bank k, diversity, and novelty,\n"
        "then compare against raw observations and a single cue-GVF baseline.",
        GREEN,
    )

    add_label(slide, left_x, 22.16, left_w, "Online GVF bank under drift", GREEN)
    add_text(
        slide,
        left_x,
        22.54,
        left_w,
        0.42,
        "A compact predictive division of labor, updated online",
        17.2,
        NAVY,
        True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_bank_mechanism(slide, left_x + 0.04, 23.06, left_w - 0.08)
    add_text(
        slide,
        left_x + 0.04,
        24.51,
        left_w - 0.08,
        0.46,
        "Mutate horizon / step size; retain candidates that fill missing roles and replenish online.",
        14.0,
        MID,
        False,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_rule(slide, left_x, 25.09, left_w, LIGHT_LINE)

    add_label(slide, left_x, 25.43, left_w, "Experimental controls", GREEN)
    add_text(
        slide,
        left_x + 0.04,
        25.85,
        left_w - 0.08,
        1.18,
        "• One-pass online learning; no replay, offline fitting, or recurrent memory.\n"
        "• Held-out probes are diagnostic only; budgets and seeds are matched within comparisons.\n"
        "• Environment-specific metrics; n=3/4 summaries are descriptive, not significance tests.",
        14.8,
        INK,
        False,
        line_spacing=0.96,
    )
    add_text(
        slide,
        left_x + 0.04,
        27.20,
        left_w - 0.08,
        0.42,
        "READOUTS   Decoding → preservation  ·  post-switch → control  ·  recovery → adaptation",
        13.5,
        MID,
        True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_label(slide, left_x, 27.85, left_w, "Why the sequence matters", BLUE)
    add_text(
        slide,
        left_x,
        28.26,
        left_w,
        1.52,
        "RQ1 tests whether feature reshaping generalizes across hidden structures.\n"
        "RQ2 fixes predictions to isolate controller accessibility. RQ3 asks how\n"
        "to compose predictions as relevance drifts. Together, the sequence moves\n"
        "from diagnosis to intervention to adaptive construction.",
        15.7,
        INK,
        False,
        line_spacing=0.98,
    )
    # ------------------------------------------------------------------
    # RIGHT COLUMN — each result answers one explicit experimental question.
    # ------------------------------------------------------------------
    add_label(slide, right_x, 9.34, right_w, "Results", GREEN)
    add_text(
        slide,
        right_x,
        9.72,
        right_w,
        0.58,
        "Three questions, three bounded answers",
        22.5,
        NAVY,
        True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_rule(slide, right_x, 10.36, right_w, GREEN, 0.035)

    add_text(
        slide,
        right_x,
        10.57,
        right_w,
        0.46,
        "RQ1 · Does reshaping predictive features reliably improve control?",
        18.4,
        NAVY,
        True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_geometry_result_table(slide, right_x, 11.12, right_w)
    add_vertical_rule(slide, right_x + 0.02, 13.40, 0.64, BLUE, 0.050)
    add_text(
        slide,
        right_x + 0.22,
        13.38,
        right_w - 0.26,
        0.72,
        "Answer: no. Effects reverse across hidden structures; feature geometry is diagnostic, not a universal\n"
        "selection rule. With n=3/condition, the evidence supports counterexamples—not a general ranking.",
        15.7,
        INK,
        False,
        valign=MSO_ANCHOR.MIDDLE,
        line_spacing=0.97,
    )
    add_rule(slide, right_x, 14.30, right_w, LIGHT_LINE)

    add_text(
        slide,
        right_x,
        14.51,
        right_w,
        0.48,
        "RQ2 · Is decodable information accessible to online control?",
        18.4,
        NAVY,
        True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        right_x,
        14.99,
        right_w,
        0.40,
        "The past cue is perfectly decodable at the junction (1.00); same gₜ and learner, only f changes.",
        15.1,
        OCHRE,
        True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    with tempfile.TemporaryDirectory(prefix="poster-figure-") as tmpdir:
        figure_path = Path(tmpdir) / "trace_only_rescue.png"
        with zipfile.ZipFile(ADAPTER_ARCHIVE) as archive:
            figure_path.write_bytes(archive.read(ADAPTER_FIGURE))
        picture = slide.shapes.add_picture(
            str(figure_path),
            inches(right_x + 0.08),
            inches(15.50),
            width=inches(right_w - 0.16),
        )
        picture.name = "Paired adapter result"
    add_text(
        slide,
        right_x + 0.08,
        22.48,
        right_w - 0.16,
        0.36,
        "Identity 0.5165  ·  RFF 0.5380  ·  Tile 0.9780  ·  Oracle / direct cue 0.9835",
        15.0,
        MID,
        False,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        right_x + 0.08,
        22.87,
        right_w - 0.16,
        0.36,
        "Tile − identity = +0.4615 ± 0.0144 SEM  ·  10/10 paired improvements",
        15.0,
        GREEN,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_vertical_rule(slide, right_x + 0.02, 23.38, 0.74, GREEN, 0.050)
    add_text(
        slide,
        right_x + 0.22,
        23.30,
        right_w - 0.26,
        0.92,
        "Answer: not necessarily. Tile coding adds no cue information; local, sparse features expose the retained\n"
        "cue to the fixed linear learner, moving control near oracle—an accessibility, not prediction, gain.",
        15.7,
        INK,
        False,
        valign=MSO_ANCHOR.MIDDLE,
        line_spacing=0.96,
    )
    add_rule(slide, right_x, 24.38, right_w, LIGHT_LINE)

    add_text(
        slide,
        right_x,
        24.58,
        right_w,
        0.48,
        "RQ3 · What bank structure converts predictions into control under drift?",
        18.4,
        NAVY,
        True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        right_x,
        25.07,
        right_w,
        0.36,
        "Raw 0.5073  ·  Cue GVF 0.5246  ·  four-seed means; uncertainty not supplied",
        14.6,
        MID,
        False,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_bank_table(slide, right_x, 25.50, right_w)
    add_text(
        slide,
        right_x + 0.10,
        27.95,
        right_w - 0.20,
        0.80,
        "• Remove diversity: −3.69 percentage points (0.5565 → 0.5196).\n"
        "• k=4 decodes more (0.7386 > 0.7294) but controls worse (0.5408 < 0.5565).\n"
        "• Remove novelty: no observed mean change (0.5565).",
        14.8,
        INK,
        False,
        line_spacing=0.98,
    )
    add_vertical_rule(slide, right_x + 0.02, 28.88, 0.74, GREEN, 0.050)
    add_text(
        slide,
        right_x + 0.22,
        28.79,
        right_w - 0.26,
        0.96,
        "Answer: supplied means favor a compact, diversity-aware bank over increasing size or novelty alone.\n"
        "The k=4 counterexample shows that more decodable predictions can still add redundancy or interference.",
        15.7,
        INK,
        False,
        valign=MSO_ANCHOR.MIDDLE,
        line_spacing=0.96,
    )

    # Full-width conclusion: mechanism statement, boundary, and next step.
    add_rule(slide, 0.92, 29.95, 21.78, NAVY, 0.055, name="Conclusion rule")
    add_label(slide, 0.92, 30.20, 3.20, "Conclusion", BLUE)
    add_text(
        slide,
        0.92,
        30.60,
        21.78,
        0.78,
        "In these streaming POMDPs, usefulness depends on the representation–learner interface—not on the predictive feature alone.",
        22.0,
        NAVY,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        line_spacing=0.96,
        name="Conclusion statement",
    )
    add_text(
        slide,
        0.92,
        31.42,
        21.78,
        0.50,
        "Same cue information, better access: tile coding moves control 0.5165 → 0.9780. More decodability, worse control: k=4 underperforms k=3—two counterexamples to selection by decoding alone.",
        15.2,
        MID,
        False,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        0.92,
        31.95,
        21.78,
        0.30,
        "Scope: controlled POMDPs · linear control · hand-structured GVF roles · n=3/4/10 by study.  Next: jointly adapt bank composition and controller coupling.",
        14.5,
        MID,
        False,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    # Footer sits directly on the original template's blue-purple gradient.
    add_rule(slide, 0.92, 32.58, 21.78, WHITE, 0.025)
    add_text(
        slide,
        0.92,
        33.46,
        6.75,
        0.34,
        "REFERENCE",
        13.8,
        WHITE,
        True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        0.92,
        33.82,
        6.75,
        0.42,
        "Sutton et al. · The Alberta Plan for AI Research · 2022",
        13.4,
        WHITE,
        False,
    )
    add_text(
        slide,
        7.87,
        33.46,
        7.90,
        0.34,
        "REPRODUCIBILITY",
        13.8,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        7.87,
        33.82,
        7.90,
        0.68,
        "Streaming · replay-free · CPU-scale\nHeld-out probes are diagnostic only",
        13.4,
        WHITE,
        False,
        align=PP_ALIGN.CENTER,
        line_spacing=0.95,
    )
    add_text(
        slide,
        15.97,
        33.46,
        6.73,
        0.34,
        "SOURCE VERSIONS",
        13.8,
        WHITE,
        True,
        align=PP_ALIGN.RIGHT,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        15.97,
        33.82,
        6.73,
        0.94,
        "Geometry 2940182f · Adapters a88813c\nGVF bank: supplied 4-seed summaries\nCommands/results accompany the course package",
        13.4,
        WHITE,
        False,
        align=PP_ALIGN.RIGHT,
        line_spacing=0.95,
    )

    presentation.save(str(OUTPUT))
    if sha256(TEMPLATE) != template_hash:
        raise RuntimeError("The source template changed during poster generation")


if __name__ == "__main__":
    build_poster()
    print(OUTPUT)
